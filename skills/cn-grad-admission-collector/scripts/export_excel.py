# -*- coding: utf-8 -*-
"""导出 Excel（可筛选明细 + 采集报告 + 冲突 + 来源日志）并生成 Markdown 采集报告。

用法：
    python export_excel.py --run ./data/考研招生计划_2027
    python export_excel.py --run <dir> --out ./输出/2027招生计划.xlsx --no-md

依赖：openpyxl
默认不覆盖已存在的输出文件，同名时自动追加时间戳。
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    CONFLICT_FIELDS, CONFLICT_HEADER_MAP, EXAM_MODES, RECORD_FIELDS,
    RECORD_HEADER_MAP, SCHOOL_FIELDS, SCHOOL_HEADER_MAP, SOURCE_FIELDS,
    SOURCE_HEADER_MAP, SOURCE_TYPES, STUDY_MODES, VERIFY_STATUS, now_str,
    read_rows, stamp, today_str,
)

QUOTA_KEYS = ("planned_total", "recommend_quota", "unified_quota")
VERIFY_FILL = {
    "待核验": "FFF2CC",
    "口径冲突": "F8CBAD",
    "无数据": "D9D9D9",
}
HEADER_FILL = "1F4E79"
HEADER_FONT = "FFFFFF"


# ---------------------------------------------------------------- 统计
def build_stats(records, schools, sources, conflicts):
    st = {}
    st["total"] = len(records)
    st["by_year"] = Counter(r.get("year", "") for r in records)
    st["by_province"] = Counter(r.get("province", "") or "未填" for r in records)
    st["by_school"] = Counter(r.get("school", "") for r in records)
    st["by_status"] = Counter(r.get("verify_status", "") or "未标注" for r in records)
    st["no_quota"] = [r for r in records if not any(r.get(k) for k in QUOTA_KEYS)]
    st["no_basis"] = [r for r in records if not r.get("quota_basis")
                      and any(r.get(k) for k in QUOTA_KEYS)]
    st["abnormal"] = []
    for r in records:
        p, rec, u = (r.get("planned_total"), r.get("recommend_quota"), r.get("unified_quota"))
        if u and p and rec and int(u) != int(p) - int(rec):
            st["abnormal"].append((r, f"统考 {u} ≠ 拟招生 {p} − 推免 {rec}"))
        elif u and not p:
            st["abnormal"].append((r, "仅有统考人数，缺拟招生人数"))
    st["need_check"] = [r for r in records if r.get("verify_status") in ("待核验", "口径冲突", "无数据")]

    empty_school_rows = [s for s in schools if not s.get("school")]
    real_schools = [s for s in schools if s.get("school")]
    st["schools_total"] = len(real_schools)
    st["schools_done"] = [s for s in real_schools if s.get("status") == "已完成"]
    st["schools_partial"] = [s for s in real_schools if s.get("status") == "部分完成"]
    st["schools_pending"] = [s for s in real_schools
                             if s.get("status") in ("待检索", "检索中", "")]
    st["schools_blocked"] = [s for s in real_schools
                             if s.get("status") in ("未完成", "访问受限")]
    st["schools_placeholder"] = empty_school_rows
    st["sources"] = sources
    st["conflicts"] = conflicts
    failed = [s for s in sources if s.get("access_result") and s.get("access_result") != "成功"]
    st["source_failed"] = failed
    return st


def md_cell(text):
    return str(text or "").replace("|", "／").replace("\n", " ")


# ---------------------------------------------------------------- Markdown 报告
def build_markdown(st, records, year, run_dir):
    L = []
    L.append(f"# {year} 年硕士研究生招生计划采集报告\n")
    L.append(f"- 生成时间：{now_str()}")
    L.append(f"- 任务目录：`{run_dir}`")
    L.append(f"- 记录总数：{st['total']} 条（涉及学校 {len(st['by_school'])} 所，省份 {len([k for k in st['by_province'] if k])} 个）\n")

    L.append("## 一、总体进度\n")
    L.append("| 指标 | 数值 |")
    L.append("| --- | --- |")
    L.append(f"| 计划采集学校 | {st['schools_total']} |")
    L.append(f"| 已完成 | {len(st['schools_done'])} |")
    L.append(f"| 部分完成 | {len(st['schools_partial'])} |")
    L.append(f"| 未开始/进行中 | {len(st['schools_pending'])} |")
    L.append(f"| 未完成/访问受限 | {len(st['schools_blocked'])} |")
    for k in VERIFY_STATUS:
        L.append(f"| 核验状态 = {k} | {st['by_status'].get(k, 0)} |")
    L.append("")

    L.append("## 二、已完成学校\n")
    if st["schools_done"]:
        L.append("| 省份 | 学校 | 代码 | 记录条数 |")
        L.append("| --- | --- | --- | --- |")
        for s in st["schools_done"]:
            L.append(f"| {md_cell(s.get('province'))} | {md_cell(s.get('school'))} | "
                     f"{md_cell(s.get('school_code'))} | {md_cell(s.get('records_count'))} |")
    else:
        L.append("（暂无）")
    L.append("")

    L.append("## 三、未完成学校\n")
    todo = st["schools_partial"] + st["schools_pending"] + st["schools_blocked"]
    if todo:
        L.append("| 省份 | 学校 | 状态 | 说明 |")
        L.append("| --- | --- | --- | --- |")
        for s in todo:
            L.append(f"| {md_cell(s.get('province'))} | {md_cell(s.get('school'))} | "
                     f"{md_cell(s.get('status'))} | {md_cell(s.get('notes'))} |")
    else:
        L.append("（无）")
    L.append("")

    L.append("## 四、访问失败与受限记录\n")
    if st["source_failed"]:
        L.append("| 学校 | 材料标题 | 访问结果 | 网址 |")
        L.append("| --- | --- | --- | --- |")
        for s in st["source_failed"]:
            L.append(f"| {md_cell(s.get('school'))} | {md_cell(s.get('source_title'))} | "
                     f"{md_cell(s.get('access_result'))} | {md_cell(s.get('url'))} |")
    else:
        L.append("（无。若实际遇到验证码/登录墙，必须在 sources.csv 登记后再导出）")
    L.append("")

    L.append("## 五、缺失数据\n")
    L.append(f"- 三项人数全部为空：{len(st['no_quota'])} 条")
    L.append(f"- 有人数但缺“人数口径说明”：{len(st['no_basis'])} 条")
    if st["no_quota"]:
        L.append("\n| 学校 | 专业代码 | 专业 | 研究方向 | 原文标题 |")
        L.append("| --- | --- | --- | --- | --- |")
        for r in st["no_quota"][:50]:
            L.append(f"| {md_cell(r.get('school'))} | {md_cell(r.get('major_code'))} | "
                     f"{md_cell(r.get('major'))} | {md_cell(r.get('direction'))} | "
                     f"{md_cell(r.get('source_title'))} |")
        if len(st["no_quota"]) > 50:
            L.append(f"\n（其余 {len(st['no_quota']) - 50} 条见 Excel「招生计划明细」筛选）")
    L.append("")

    L.append("## 六、异常数据\n")
    if st["abnormal"]:
        L.append("| 学校 | 专业代码 | 异常说明 | 原文标题 |")
        L.append("| --- | --- | --- | --- |")
        for r, why in st["abnormal"]:
            L.append(f"| {md_cell(r.get('school'))} | {md_cell(r.get('major_code'))} | "
                     f"{md_cell(why)} | {md_cell(r.get('source_title'))} |")
    else:
        L.append("（无数值关系异常）")
    L.append("")

    L.append("## 七、待人工核验项目\n")
    if st["need_check"]:
        L.append(f"共 {len(st['need_check'])} 条。")
        L.append("\n| 核验状态 | 学校 | 专业代码 | 研究方向 | 来源 |")
        L.append("| --- | --- | --- | --- | --- |")
        for r in st["need_check"][:60]:
            L.append(f"| {md_cell(r.get('verify_status'))} | {md_cell(r.get('school'))} | "
                     f"{md_cell(r.get('major_code'))} | {md_cell(r.get('direction'))} | "
                     f"{md_cell(r.get('source_url'))} |")
        if len(st["need_check"]) > 60:
            L.append(f"\n（其余 {len(st['need_check']) - 60} 条见 Excel 筛选）")
    else:
        L.append("（无）")
    L.append("")

    L.append("## 八、口径冲突\n")
    if st["conflicts"]:
        L.append("| 学校 | 专业代码 | 字段 | 来源A值 | 来源A | 来源B值 | 来源B | 处理 |")
        L.append("| --- | --- | --- | --- | --- | --- | --- | --- |")
        for c in st["conflicts"]:
            L.append(f"| {md_cell(c.get('school'))} | {md_cell(c.get('major_code'))} | "
                     f"{md_cell(c.get('field'))} | {md_cell(c.get('value_a'))} | "
                     f"{md_cell(c.get('source_a'))} | {md_cell(c.get('value_b'))} | "
                     f"{md_cell(c.get('source_b'))} | {md_cell(c.get('resolution'))} |")
    else:
        L.append("（未检测到跨来源数值差异）")
    L.append("")

    L.append("## 九、来源统计\n")
    L.append(f"- 登记来源 {len(st['sources'])} 条，其中访问失败 {len(st['source_failed'])} 条")
    types = Counter(s.get("source_type", "") or "未标注" for s in st["sources"])
    for k, v in types.most_common():
        L.append(f"- 材料类型「{k}」：{v} 条")
    L.append("")
    L.append("> 说明：本报告所有人数均来自 records.csv 中登记的官方原文；「待核验」项禁止作为最终结论引用。")
    return "\n".join(L)


# ---------------------------------------------------------------- Excel
def style_header(ws, ncols, row=1):
    from openpyxl.styles import Alignment, Font, PatternFill
    fill = PatternFill("solid", fgColor=HEADER_FILL)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = Font(color=HEADER_FONT, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def write_table(ws, fields, rows, widths=True, start_row=1, numeric_keys=()):
    from openpyxl.utils import get_column_letter
    headers = [h for _k, h, _w in fields]
    keys = [k for k, _h, _w in fields]
    for j, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=j, value=h)
    for i, row in enumerate(rows, start_row + 1):
        for j, k in enumerate(keys, 1):
            v = row.get(k, "")
            if v != "" and k in numeric_keys and str(v).isdigit():
                v = int(v)
            ws.cell(row=i, column=j, value=v if v != "" else None)
    style_header(ws, len(headers), start_row)
    if widths:
        for j, (_k, _h, w) in enumerate(fields, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + len(rows)


def export_excel(run_dir, out_path, st, records, schools, sources, conflicts, year):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # ---- Sheet1 明细
    ws = wb.active
    ws.title = "招生计划明细"
    last = write_table(ws, RECORD_FIELDS, records,
                       numeric_keys=("year",) + QUOTA_KEYS)
    ncol = len(RECORD_FIELDS)
    ref = f"A1:{get_column_letter(ncol)}{max(last, 2)}"
    ws.auto_filter.ref = ref
    ws.freeze_panes = "B2"

    col = {k: get_column_letter(i) for i, (k, _h, _w) in enumerate(RECORD_FIELDS, 1)}
    # 状态列条件格式
    status_ref = f"{col['verify_status']}2:{col['verify_status']}{max(last, 2) + 1}"
    for status, color in VERIFY_FILL.items():
        ws.conditional_formatting.add(
            status_ref,
            FormulaRule(formula=[f'${col["verify_status"]}2="{status}"'],
                        fill=PatternFill("solid", bgColor=color), stopIfTrue=False),
        )
    # 三项人数全空 → 灰色提示
    if last >= 2:
        ws.conditional_formatting.add(
            f"{col['planned_total']}2:{col['planned_total']}{last}",
            FormulaRule(formula=[
                f'AND(${col["planned_total"]}2="",${col["recommend_quota"]}2="",${col["unified_quota"]}2="")'
            ], fill=PatternFill("solid", bgColor="D9D9D9")),
        )

    # 下拉枚举
    dv_map = {
        "verify_status": VERIFY_STATUS, "study_mode": STUDY_MODES,
        "exam_mode": EXAM_MODES, "source_type": SOURCE_TYPES,
    }
    end_row = max(last, 2) + 500
    for key, options in dv_map.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                            allow_blank=True, showDropDown=False)
        dv.error = "请选择枚举内取值，或按规则新增强调说明"
        ws.add_data_validation(dv)
        dv.add(f"{col[key]}2:{col[key]}{end_row}")

    # 超链接 + 长文本换行 + 数字列格式
    for r in range(2, last + 1):
        u = ws.cell(row=r, column=list(col).index("source_url") + 1)
        if u.value and str(u.value).startswith("http"):
            u.hyperlink = str(u.value)
            u.font = Font(color="0563C1", underline="single", size=10)
        for key in ("quota_basis", "notes", "direction", "source_title"):
            ws.cell(row=r, column=list(col).index(key) + 1).alignment = Alignment(
                vertical="top", wrap_text=True)
        for key in ("school_code", "major_code", "college_code"):
            ws.cell(row=r, column=list(col).index(key) + 1).number_format = "@"
    ws.sheet_view.zoomScale = 95

    # ---- Sheet2 采集报告
    ws2 = wb.create_sheet("采集报告")
    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 22
    ws2.column_dimensions["C"].width = 60
    ws2["A1"] = f"{year} 年硕士研究生招生计划 · 采集报告"
    ws2["A1"].font = Font(bold=True, size=14)
    rows2 = [
        ("生成时间", now_str(), ""),
        ("记录总数", st["total"], "每条均可回溯到登记来源"),
        ("涉及学校数", len(st["by_school"]), ""),
        ("计划采集学校", st["schools_total"], ""),
        ("已完成学校", len(st["schools_done"]), ""),
        ("部分完成", len(st["schools_partial"]), ""),
        ("未开始/进行中", len(st["schools_pending"]), ""),
        ("未完成/访问受限", len(st["schools_blocked"]), ""),
        ("待核验记录", st["by_status"].get("待核验", 0), "禁止直接引用为最终人数"),
        ("口径冲突记录", st["by_status"].get("口径冲突", 0), "见「口径冲突」表"),
        ("无数据记录", st["by_status"].get("无数据", 0), "官方材料未列明人数"),
        ("三项人数全空", len(st["no_quota"]), ""),
        ("缺人数口径说明", len(st["no_basis"]), "不得据此认定人数口径"),
        ("数值关系异常", len(st["abnormal"]), "统考 ≠ 拟招生 − 推免 等"),
        ("来源登记", len(sources), f"其中访问失败 {len(st['source_failed'])} 条"),
    ]
    r = 3
    for label, val, note in rows2:
        ws2.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws2.cell(row=r, column=2, value=val)
        ws2.cell(row=r, column=3, value=note)
        r += 1
    r += 1
    ws2.cell(row=r, column=1, value="数据边界声明").font = Font(bold=True, color="C00000")
    ws2.cell(row=r, column=3, value="人数仅取自官方原文（研究生院/研招网/学院官网/研招网专业目录）；"
                                   "搜索摘要、机构与往年资料不得作为当年人数依据；不推算统考名额。")
    ws2.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    ws2.row_dimensions[r].height = 48

    # ---- Sheet3 口径冲突
    ws3 = wb.create_sheet("口径冲突")
    write_table(ws3, CONFLICT_FIELDS, conflicts)
    if conflicts:
        ws3.auto_filter.ref = f"A1:{get_column_letter(len(CONFLICT_FIELDS))}{len(conflicts) + 1}"
    ws3.freeze_panes = "A2"

    # ---- Sheet4 来源日志
    ws4 = wb.create_sheet("来源日志")
    write_table(ws4, SOURCE_FIELDS, sources)
    if sources:
        ws4.auto_filter.ref = f"A1:{get_column_letter(len(SOURCE_FIELDS))}{len(sources) + 1}"
    ws4.freeze_panes = "A2"
    for r in range(2, len(sources) + 2):
        c = ws4.cell(row=r, column=6)
        if c.value and str(c.value).startswith("http"):
            c.hyperlink = str(c.value)
            c.font = Font(color="0563C1", underline="single", size=10)

    # ---- Sheet5 学校进度
    ws5 = wb.create_sheet("学校清单")
    write_table(ws5, SCHOOL_FIELDS, schools)
    if schools:
        ws5.auto_filter.ref = f"A1:{get_column_letter(len(SCHOOL_FIELDS))}{len(schools) + 1}"
    ws5.freeze_panes = "A2"

    # ---- Sheet6 字段说明
    ws6 = wb.create_sheet("字段说明")
    ws6.column_dimensions["A"].width = 16
    ws6.column_dimensions["B"].width = 18
    ws6.column_dimensions["C"].width = 78
    ws6["A1"] = "字段"
    ws6["B1"] = "中文名"
    ws6["C1"] = "填写规则"
    style_header(ws6, 3)
    rules = {
        "year": "招生年份，四位数字；以招生简章标题/发布时间归属年份为准",
        "school_code": "教育部五位数院校代码，取自研招网或简章封面，不得猜测",
        "major_code": "六位专业代码（可加两位方向码），如 085400 / 085400-01",
        "study_mode": "全日制 / 非全日制 / 全日制、非全日制分别招生时分行填写",
        "exam_mode": "全国统考 / 推荐免试 / 单独考试 / 管理类联考 等，以原文表述为准",
        "planned_total": "拟招生人数：仅当原文明确写出该专业（方向）拟招总数时填写",
        "recommend_quota": "推免人数：仅当原文明确列出推免/拟接收推免人数时填写",
        "unified_quota": "统考计划人数：仅当原文明确单列统考名额时填写；禁止用拟招生减推免推算",
        "quota_basis": "抄录原文对人数口径的表述（是否含推免、是否含非全、是否学院口径）",
        "source_title": "官方原文标题，与网页/PDF 标题一致",
        "source_url": "可追溯的原始链接（.edu.cn / chsi.com.cn / 官方 PDF），不用搜索跳转链接",
        "publish_date": "原文发布日期 YYYY-MM-DD；PDF 取文首日期或网页发布时间",
        "verify_status": "已核验 / 待核验 / 口径冲突 / 无数据",
        "is_current": "多版本并存时：当前采用 / 历史版本（历史记录不删除）",
    }
    r = 2
    for k, h, _w in RECORD_FIELDS:
        ws6.cell(row=r, column=1, value=k)
        ws6.cell(row=r, column=2, value=h)
        ws6.cell(row=r, column=3, value=rules.get(k, "")).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description="导出招生计划 Excel 与采集报告")
    ap.add_argument("--run", required=True, help="任务根目录")
    ap.add_argument("--out", default="", help="xlsx 输出路径")
    ap.add_argument("--no-md", action="store_true", help="不生成 Markdown 报告")
    args = ap.parse_args()

    run = Path(args.run).expanduser().resolve()
    records = read_rows(run / "data" / "records.csv", RECORD_HEADER_MAP)
    if not records:
        print(f"[错误] {run/'data'/'records.csv'} 无数据或不存在", file=sys.stderr)
        sys.exit(2)
    schools = read_rows(run / "data" / "schools.csv", SCHOOL_HEADER_MAP)
    sources = read_rows(run / "data" / "sources.csv", SOURCE_HEADER_MAP)
    conflicts = read_rows(run / "data" / "conflicts.csv", CONFLICT_HEADER_MAP)

    years = Counter(r.get("year", "") for r in records)
    year = years.most_common(1)[0][0] or today_str()[:4]

    out = Path(args.out).expanduser() if args.out else run / f"考研招生计划_{year}.xlsx"
    if not out.is_absolute():
        out = (Path.cwd() / out).resolve()
    if out.exists():
        out = out.with_name(f"{out.stem}_{stamp()}{out.suffix}")
    out.parent.mkdir(parents=True, exist_ok=True)

    st = build_stats(records, schools, sources, conflicts)
    export_excel(run, out, st, records, schools, sources, conflicts, year)
    print(f"[Excel] {out}")

    if not args.no_md:
        md_path = run / f"采集报告_{year}.md"
        if md_path.exists():
            md_path = md_path.with_name(f"采集报告_{year}_{stamp()}.md")
        md_path.write_text(build_markdown(st, records, year, run), encoding="utf-8")
        print(f"[报告]  {md_path}")

    print(f"[统计] 记录 {st['total']} 条｜待核验 {st['by_status'].get('待核验', 0)}｜"
          f"口径冲突 {st['by_status'].get('口径冲突', 0)}｜无数值 {len(st['no_quota'])}｜"
          f"异常 {len(st['abnormal'])}")


if __name__ == "__main__":
    main()
